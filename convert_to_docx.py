#!/usr/bin/env python3
"""
Convert Markdown documentation to Word (.docx) format
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import re

def add_heading_style(doc, text, level=1):
    """Add heading with proper styling"""
    if level == 1:
        doc.add_heading(text, level=1)
    elif level == 2:
        doc.add_heading(text, level=2)
    else:
        doc.add_heading(text, level=3)

def add_styled_paragraph(doc, text, style='Normal', bold=False, italic=False, color=None):
    """Add paragraph with styling"""
    p = doc.add_paragraph(text, style=style)
    if bold or italic or color:
        for run in p.runs:
            if bold:
                run.bold = True
            if italic:
                run.italic = True
            if color:
                run.font.color.rgb = color
    return p

def convert_markdown_to_docx():
    """Convert markdown file to Word document"""
    
    # Read markdown file
    with open('COMPLETE_PROJECT_DOCUMENTATION.md', 'r', encoding='utf-8') as f:
        content = f.read()
    
    # Create Word document
    doc = Document()
    
    # Set document margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
    
    # Add title page
    title = doc.add_heading('COMPLETE PROJECT DOCUMENTATION', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    subtitle = doc.add_paragraph('Online Food Ordering Web Application')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.runs[0]
    subtitle_run.bold = True
    subtitle_run.font.size = Pt(14)
    
    # Add metadata
    doc.add_paragraph('')
    metadata = doc.add_paragraph()
    metadata.add_run('Document Version: ').bold = True
    metadata.add_run('1.0\n')
    
    metadata = doc.add_paragraph()
    metadata.add_run('Date: ').bold = True
    metadata.add_run('June 24, 2026\n')
    
    metadata = doc.add_paragraph()
    metadata.add_run('Status: ').bold = True
    metadata.add_run('✅ Complete & Production Ready\n')
    
    metadata = doc.add_paragraph()
    metadata.add_run('Team: ').bold = True
    metadata.add_run('Full-Stack Development\n')
    
    # Add page break
    doc.add_page_break()
    
    # Process content line by line
    lines = content.split('\n')
    i = 0
    
    while i < len(lines):
        line = lines[i]
        
        # Skip empty lines
        if not line.strip():
            i += 1
            continue
        
        # Handle headings
        if line.startswith('# '):
            text = line.replace('# ', '').strip()
            add_heading_style(doc, text, level=1)
        elif line.startswith('## '):
            text = line.replace('## ', '').strip()
            add_heading_style(doc, text, level=2)
        elif line.startswith('### '):
            text = line.replace('### ', '').strip()
            add_heading_style(doc, text, level=3)
        
        # Handle code blocks
        elif line.strip().startswith('```'):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            
            # Add code as formatted text
            code_text = '\n'.join(code_lines)
            if code_text.strip():
                p = doc.add_paragraph(code_text)
                p.style = 'List Bullet'
                for run in p.runs:
                    run.font.name = 'Courier New'
                    run.font.size = Pt(10)
        
        # Handle tables
        elif '|' in line and i + 1 < len(lines) and '|' in lines[i + 1]:
            # Parse table
            table_rows = []
            while i < len(lines) and '|' in lines[i]:
                cells = [cell.strip() for cell in lines[i].split('|')]
                cells = [c for c in cells if c and c != '-' * len(c)]  # Remove separators
                if cells:
                    table_rows.append(cells)
                i += 1
            
            if table_rows:
                # Create table
                table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
                table.style = 'Light Grid Accent 1'
                
                for row_idx, row_data in enumerate(table_rows):
                    for col_idx, cell_data in enumerate(row_data):
                        cell = table.rows[row_idx].cells[col_idx]
                        cell.text = cell_data
                        
                        # Style header row
                        if row_idx == 0:
                            for paragraph in cell.paragraphs:
                                for run in paragraph.runs:
                                    run.bold = True
            i -= 1
        
        # Handle bullet points
        elif line.strip().startswith('- '):
            text = line.replace('- ', '').strip()
            doc.add_paragraph(text, style='List Bullet')
        
        # Handle numbered lists
        elif re.match(r'^\d+\. ', line.strip()):
            text = re.sub(r'^\d+\. ', '', line.strip())
            doc.add_paragraph(text, style='List Number')
        
        # Handle bold text (markdown style)
        elif '**' in line:
            p = doc.add_paragraph()
            parts = line.split('**')
            for idx, part in enumerate(parts):
                if idx % 2 == 0:
                    p.add_run(part)
                else:
                    run = p.add_run(part)
                    run.bold = True
        
        # Handle regular paragraphs
        elif line.strip():
            doc.add_paragraph(line.strip())
        
        i += 1
    
    # Add footer
    doc.add_page_break()
    footer_heading = doc.add_heading('Document Information', level=2)
    
    footer_para = doc.add_paragraph()
    footer_para.add_run('Version: ').bold = True
    footer_para.add_run('1.0\n')
    
    footer_para = doc.add_paragraph()
    footer_para.add_run('Created: ').bold = True
    footer_para.add_run('June 24, 2026\n')
    
    footer_para = doc.add_paragraph()
    footer_para.add_run('Status: ').bold = True
    footer_para.add_run('Complete\n')
    
    footer_para = doc.add_paragraph()
    footer_para.add_run('Language: ').bold = True
    footer_para.add_run('English\n')
    
    footer_para = doc.add_paragraph()
    footer_para.add_run('Distribution: ').bold = True
    footer_para.add_run('Development Team, Product Manager, DevOps Team, Client\n')
    
    # Save document
    doc.save('COMPLETE_PROJECT_DOCUMENTATION.docx')
    print("✅ Word document created successfully!")
    print("📄 File: COMPLETE_PROJECT_DOCUMENTATION.docx")

if __name__ == '__main__':
    convert_markdown_to_docx()
