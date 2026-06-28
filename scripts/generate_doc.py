from docx import Document
from docx.shared import Pt
import os

project_root = os.path.dirname(os.path.dirname(__file__))

def add_heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def add_paragraph(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.bold = bold


def main():
    doc = Document()
    doc.core_properties.title = 'Online Food Ordering - Project Documentation'

    add_heading(doc, 'Project Overview', level=1)
    add_paragraph(doc, 'This project is an online food ordering web application built with a Node.js + Express backend and a React frontend (Vite).')

    add_heading(doc, 'Repository Structure', level=1)
    add_paragraph(doc, 'Top-level files and folders:')
    structure = [
        'backend/ - Node.js + Express backend with routes, models, and middleware',
        'frontend/ - React frontend built with Vite, components, pages, and styles',
        'README.md, QUICKSTART.md and various deployment and setup documentation files',
    ]
    for s in structure:
        add_paragraph(doc, f'- {s}')

    add_heading(doc, 'Backend', level=1)
    add_paragraph(doc, 'The backend is implemented using Node.js and Express. Key files:')
    backend_items = [
        'server.js - Express server entrypoint',
        'routes/ - API endpoints for auth, products, cart, orders, users, admin',
        'models/ - Mongoose models for User, Product, Order',
        'middleware/ - auth middleware',
    ]
    for item in backend_items:
        add_paragraph(doc, f'- {item}')

    add_heading(doc, 'Frontend', level=1)
    add_paragraph(doc, 'The frontend is a React application with Vite. Key parts:')
    frontend_items = [
        'src/pages/ - Pages like Home, Menu, Cart, Login, Register, Profile, AdminDashboard',
        'src/components/ - Reusable components like Navbar, ProductCard, CartItem',
        'src/styles/ - CSS files for each page and component',
        'src/store/useStore.js - simple global state store',
        'i18n/ - localization configuration and locales',
    ]
    for item in frontend_items:
        add_paragraph(doc, f'- {item}')

    add_heading(doc, 'Setup and Quickstart', level=1)
    add_paragraph(doc, 'Backend:')
    add_paragraph(doc, '1) Install dependencies: `npm install` in the backend folder')
    add_paragraph(doc, '2) Start the backend: `node server.js` or use a process manager like PM2')
    add_paragraph(doc, 'Frontend:')
    add_paragraph(doc, '1) Install dependencies: `npm install` in the frontend folder')
    add_paragraph(doc, '2) Dev server: `npm run dev` (Vite). Build: `npm run build`. Serve the `dist` folder with any static server, or the backend as configured.')

    add_heading(doc, 'Deployment', level=1)
    add_paragraph(doc, 'The project includes Dockerfiles for both frontend and backend. You can build and push images to a registry and deploy using Docker Compose or a container platform like Railway, Heroku, or AWS ECS.')

    add_heading(doc, 'Database', level=1)
    add_paragraph(doc, 'The application uses MongoDB. Connection string is read from `MONGODB_URI` environment variable. Seed script `seedDatabase.js` is provided to populate sample data.')

    add_heading(doc, 'Important Files', level=1)
    important_files = [
        'backend/server.js',
        'backend/routes/*.js',
        'frontend/src/App.jsx',
        'frontend/src/pages/Home.jsx',
        'frontend/src/styles/Home.css',
    ]
    for f in important_files:
        add_paragraph(doc, f'- {f}')

    add_heading(doc, 'Troubleshooting & Notes', level=1)
    add_paragraph(doc, '1) Ensure static files are served correctly and that paths in HTML/CSS use absolute `/` paths when routing is used on the client side.')
    add_paragraph(doc, '2) Case sensitivity: filenames on production are case-sensitive (Linux).')
    add_paragraph(doc, '3) If CSS disappears in production: check missing closing braces or media query scopes in CSS files, and confirm build step includes all assets.')

    add_heading(doc, 'Contact & Maintainers', level=1)
    add_paragraph(doc, 'Maintainer: Mohamed Mostafa')
    add_paragraph(doc, 'Repository: https://github.com/mohamedmostaf21/online-food-ordering-website')

    out_path = os.path.join(project_root, 'PROJECT_DOCUMENTATION.docx')
    doc.save(out_path)
    print('Saved documentation to', out_path)

if __name__ == '__main__':
    main()
